"""Text extraction for uploaded files: TXT, VTT, SRT, PDF, DOCX (and MD as plain text)."""
from __future__ import annotations

import io
import re

SUPPORTED_EXTENSIONS = {".txt", ".vtt", ".srt", ".pdf", ".docx", ".md"}

_TIMESTAMP = re.compile(
    r"^\s*\d{1,2}:\d{2}(?::\d{2})?[.,]\d{1,3}\s*-->\s*\d{1,2}:\d{2}(?::\d{2})?[.,]\d{1,3}.*$"
)
_VOICE_TAG = re.compile(r"<v(?:\.[^\s>]+)?\s+([^>]+)>(.*?)(?:</v>|$)", re.S)
_ANY_TAG = re.compile(r"</?[^>]+>")


class ExtractionError(Exception):
    pass


def extension_of(filename: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", filename or "")
    return m.group(1).lower() if m else ""


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")


def _parse_cues(text: str) -> str:
    """Parse WebVTT / SRT into 'Speaker: utterance' lines, merging consecutive same-speaker cues."""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    out: list[tuple[str, str]] = []  # (speaker, text)
    in_note = False
    for raw in lines:
        line = raw.strip()
        if not line:
            in_note = False
            continue
        if line.startswith("WEBVTT") or line.startswith("Kind:") or line.startswith("Language:"):
            continue
        if line.startswith("NOTE") or line.startswith("STYLE") or line.startswith("REGION"):
            in_note = True
            continue
        if in_note:
            continue
        if _TIMESTAMP.match(line):
            continue
        if re.fullmatch(r"\d+", line):  # cue number
            continue
        if re.fullmatch(r"[0-9a-fA-F-]{8,}(/\d+-\d+)?", line):  # Teams-style cue ids
            continue
        speaker = ""
        m = _VOICE_TAG.search(line)
        if m:
            speaker = m.group(1).strip()
            content = _ANY_TAG.sub("", m.group(2)).strip()
        else:
            content = _ANY_TAG.sub("", line).strip()
            sm = re.match(r"^([A-Z][\w .'\-]{0,40}):\s+(.*)$", content)
            if sm:
                speaker, content = sm.group(1).strip(), sm.group(2).strip()
        if not content:
            continue
        if out and (speaker == out[-1][0] or (not speaker and out)):
            prev_speaker = out[-1][0]
            out[-1] = (prev_speaker, out[-1][1] + " " + content)
        else:
            out.append((speaker, content))
    return "\n".join(f"{s}: {t}" if s else t for s, t in out)


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # pragma: no cover - damaged page
            continue
    text = "\n".join(parts).strip()
    if not text:
        raise ExtractionError("PDF has no extractable text (it may be a scanned image).")
    return text


def _docx_text(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ExtractionError("DOCX contains no text.")
    return text


def extract_text(filename: str, data: bytes) -> str:
    ext = extension_of(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            f"Unsupported file type '{ext or 'none'}'. Use PDF, TXT, VTT, SRT or DOCX."
        )
    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    else:
        text = _decode(data)
        if ext in (".vtt", ".srt") or text.lstrip().startswith("WEBVTT"):
            text = _parse_cues(text)
    text = text.replace("\x00", "").strip()
    if not text:
        raise ExtractionError("File is empty after text extraction.")
    return text
